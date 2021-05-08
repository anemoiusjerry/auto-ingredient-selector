import os
import traceback, sys
import logging
import pandas as pd
import jinja2
import pdfkit
import json
import copy
import docx
import argparse
import math

from PySide2.QtWidgets import *
from PySide2.QtCore import QObject, Signal, Slot

from datetime import *
from openpyxl import load_workbook
from pathlib import Path
from config.configParser import FigMe

class InfoSheetGenerator(QObject):
    launched = Signal(int)
    stateChanged = Signal(int)
    cancel = Signal()
    error = Signal(str)

    def __init__(self, infoSheet_df, gdriveObject, config):
        QObject.__init__(self)
        self.stop = False

        self.config = config
        self.contents_df = infoSheet_df
        self.gdriveObject = gdriveObject
        self.errorStr = ""
        self.SOF = 8

        # Open html template as string
        if getattr(sys, 'frozen', False):
            app_path = sys._MEIPASS
            path = os.path.dirname(sys.executable)
            self.footer_save_path = os.path.abspath(os.path.join(path, os.pardir)) + "/Resources"
        else:
            app_path = os.getcwd()
            self.footer_save_path = app_path
        self.app_path = app_path

        # Normal template
        tmpPath =  app_path + "/Assets/InfoSheetTemplate.html"
        html_tmp = open(tmpPath, 'r')
        self.template = jinja2.Environment(loader=jinja2.BaseLoader).from_string(html_tmp.read())

        # Men template
        menPath = app_path + "/Assets/CourageousTemplate.html"
        cour_tmp = open(menPath, 'r')
        self.courTemplate = jinja2.Environment(loader=jinja2.BaseLoader).from_string(cour_tmp.read())

        # Allow use of len method in html
        self.template.globals["len"] = len
        self.courTemplate.globals["len"] = len

        wkhtml_path = app_path + "/wkhtmltopdf.exe"
        self.pdfkitConfig = pdfkit.configuration(wkhtmltopdf=wkhtml_path)
        self.options = {
            "orientation":"Landscape",
            "enable-local-file-access":None,
        }
        self.app_path = app_path

    def genHeader(self):
        header_tmp = open(self.app_path + "/Assets/headerTmp.html")
        headerTemplate = jinja2.Environment(loader=jinja2.BaseLoader).from_string(header_tmp.read())
        html_str = headerTemplate.render(misc_vals=self.misc_values)
        with open(f"{self.footer_save_path}/header.html", "w") as f:
            f.write(html_str)

    def genFooter(self):
        footer_tmp = open(self.app_path + "/Assets/footerTmp.html")
        footerTemplate = jinja2.Environment(loader=jinja2.BaseLoader).from_string(footer_tmp.read())
        html_str = footerTemplate.render(misc_vals=self.misc_values)
        with open(f"{self.footer_save_path}/footer.html", "w") as f:
            f.write(html_str)

    def isExcel(self, filename):
        excelLike = ["csv", "xlsx", "xls"]
        ext = filename.split(".")[-1]
        if ext in (excelLike):
            return True
        return False

    def process_all(self):
        # Get all formulation sheets from output dir
        path = self.config.getDir("Export Directory")

        # Pass all non-order dependent information
        self.misc_values = {}
        self.misc_values["assetsPath"] = self.app_path + "/Assets"
        self.misc_values["address"] = self.config.getMisc("Address")
        self.misc_values["website"] = self.config.getMisc("Website")
        self.misc_values["font"] = self.config.getMisc("Font")
        self.misc_values["contactNumber"] = self.config.getMisc("Contact number")
        self.misc_values["abn"] = self.config.getMisc("ABN")
        self.misc_values["para_offset"] = 0

        # Generate Footer html
        self.genFooter()
        self.genHeader()
        self.options["footer-html"] = self.footer_save_path + "/footer.html"
        self.options["header-html"] = self.footer_save_path + "/header.html"

        # Create list of all formlation sheet files
        sheet_paths = []
        # Go through all order folders
        for _folder in os.listdir(path):
            _folder_path = os.path.join(path, _folder)
            # Check if file is a folder
            if os.path.isdir(_folder_path):
                # Go through all files in folder
                for _file in os.listdir(_folder_path):
                    _file_path = os.path.join(_folder_path, _file)
                    if os.path.isfile(_file_path) and "Worksheet" in _file and self.isExcel(_file_path):
                        sheet_paths.append(_file_path)
            # Add to paths if is excel worksheet
            elif "Worksheet" in _folder_path and self.isExcel(_folder_path):
                sheet_paths.append(_folder_path)
        
        self.launched.emit(len(sheet_paths))
        # Retrieve all info. needed for pdf
        for i, f in enumerate(sheet_paths):
            # Stop if cancelled
            if self.stop:
                return
            self.stateChanged.emit(i)

            workbook = load_workbook(filename=f)
            sheet = workbook.active

            name = sheet["B1"].value
            prod_type = sheet["B2"].value
            isMale = False
            if "[M]" in sheet["A2"].value:
                isMale = True

            df = self.extract_incis(sheet, self.contents_df)
            df = self.fill_dates(sheet, df)
            df = self.fill_instructions(name, prod_type, df)
            headings, paragraphs = self.split_sections(df)

            # Pass all info to gen brochure
            print("calling generate report")
            try:
                self.generateReport(headings, paragraphs, name, prod_type, isMale)
            except Exception as e:
                print(e)
                self.errorStr += f"Failed to generate {name.upper()}'s {prod_type.upper()} report PDF.\n\n"
        
        if self.errorStr != "":
            self.error.emit(self.errorStr)

    def fill_instructions(self, name, prod_type, df):
        # Reset para offset
        self.misc_values["para_offset"] = 0
        instructions_filename = f"{prod_type.title()} Instructions"

        # Get product instructions
        try:
            if (self.config.masterDict["gdrive"]):
                error_msg = f"Failed to fetch {prod_type} instructions from Google Drive."
                f, file_id = self.gdriveObject.fetch_file(instructions_filename)
            else:
                error_msg = f"Failed to fetch {prod_type} instructions."
                instructions_path = self.config.getDir("Product Instructions Directory") + f"/{instructions_filename}.docx"
                f = open(instructions_path, "rb")
        # If failed to open instructions then return straightaway
        except:
            self.errorStr += error_msg
            return df

        # First name only
        fullText = name.split(" ")[0] + ", "
        doc = docx.Document(f)
        for para in doc.paragraphs:
            fullText += para.text + "\n"

        # Find word count of instructions to determine offset number
        word_count = len(fullText.split(" "))
        if word_count > 200:
            self.misc_values["para_offset"] = 1
        df[["Recommendations For Use"]] = fullText
        return df

    def fill_dates(self, sheet, df):
        # Convert date strings to datetime object
        try:

            date_blended = datetime.strptime(sheet["B3"].value, "%d.%m.%Y")
            expiry_date = datetime.strptime(sheet["B6"].value, "%d.%m.%Y")
            dt = expiry_date - date_blended
        # Return df if exception to allow for manual user editing
        except:
            name = sheet["B1"].value.upper()
            prod_type = sheet["B2"].value.upper()
            # Add to error string
            self.errorStr += f"- Date not in the format of dd.mm.yyyy for {name}'s {prod_type} order.\n\n"
            return df
        # Always round down to nearest month
        months = math.floor(dt.days / 30)
        text = df.loc[0]["Used By & Best Before Date"]
        df[["Used By & Best Before Date"]] = text.replace("...", f" {months} ")

        date_blended_str = datetime.strftime(date_blended, "%d/%m/%Y")
        expiry_date_str = datetime.strftime(expiry_date, "%d/%m/%Y")
        df[["Used By & Best Before Date"]] += f"\n\nDate Blended: {date_blended_str}\
                                                \nBest Before Date: {expiry_date_str}"
        return df

    def extract_incis(self, sheet, df):
        """ Returns df with ingredients INCI section inserted
        """
        # make an deep copy of df
        df = copy.deepcopy(df)

        inci_str = "Batch No. " + str(sheet["B4"].value) + "\n\n" + "Ingredients: "
        i=self.SOF

        inci_names = []
        # Add inci names and weights to list of tuples for sorting
        while sheet[f"C{i}"].value != None:
            inci = sheet[f"A{i}"].value
            weight = sheet[f"D{i}"].value
            if inci != None:
                inci_names.append((inci, weight))
            i+=1

        # Sort ingredients by weight (2nd element in tuple)
        inci_names.sort(key=lambda x: x[1], reverse=True)

        for inci, weight in inci_names:
            inci_str += inci + ", "
        # Remove final comma with full stop
        inci_str = inci_str[:-2] + "."

        df[["Ingredients"]] = inci_str
        return df

    def generateReport(self, headings, paragraphs, name, prod_type, isMale):
        name = name.title()
        prod_type = prod_type.title()

        # Use male template if a male product
        if isMale:
            self.misc_values["isMale"] = True
            self.genHeader()
            html_str = self.courTemplate.render(headings=headings, paragraphs=paragraphs,
                name=name, prod_type=prod_type, misc_vals=self.misc_values)
        else:
            self.misc_values["isMale"] = False
            self.genHeader()
            html_str = self.template.render(headings=headings, paragraphs=paragraphs,
                name=name, prod_type=prod_type, misc_vals=self.misc_values)

        output_path = self.config.getDir("Export Directory") + f"/{name}"

        # Create reports folder if it doesnt exist
        if not os.path.exists(output_path):
            os.makedirs(output_path)

        pdfkit.from_string(html_str,
            output_path + f"/{name} - {prod_type} - PIIS.pdf",
            configuration=self.pdfkitConfig, options=self.options)

    def split_sections(self, df):
        headings = list(df)
        paragraphs = []
        for i in range(df.shape[1]):
            paragraphs.append(df.iloc[0,i])
        return headings, paragraphs

    @Slot()
    def stop_(self):
        self.stop = True
        self.cancel.emit()